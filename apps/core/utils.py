from io import BytesIO
import json

from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4

from docx import Document

def generate_profile_pdf(content: str) -> BytesIO:
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()
    story = []

    for line in content.split("\n"):
        story.append(Paragraph(line.replace("&", "&amp;"), styles["Normal"]))

    doc.build(story)
    buffer.seek(0)
    return buffer


def generate_profile_docx(content: str) -> BytesIO:
    buffer = BytesIO()
    document = Document()

    document.add_heading("Profile Evaluation Report", level=1)

    try:
        parsed = json.loads(content)
        pretty = json.dumps(parsed, indent=2)
        document.add_paragraph(pretty)
    except Exception:
        document.add_paragraph(content)

    document.save(buffer)
    buffer.seek(0)
    return buffer
